from functools import lru_cache
import os

import json
import chardet
import pandas as pd
import time

from dotenv import load_dotenv
from pypdf import PdfReader
from docx2txt import process as docx_process
from pptx import Presentation

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma
from langchain_community.embeddings import HuggingFaceEmbeddings

# ============================================================
# ENVIRONMENT
# ============================================================

load_dotenv(".env")

CHROMA_PATH = "./chroma_db_robust"

# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".xlsx",
    ".xls",
    ".json",
    ".pptx",
}


# ============================================================
# EMBEDDING MODEL
# ============================================================

EMBEDDING_MODEL = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"

_embeddings = None


def get_embeddings():
    """
    Create the embedding model once and reuse it.
    """
    global _embeddings

    if _embeddings is None:
        _embeddings = HuggingFaceEmbeddings(
            model_name=EMBEDDING_MODEL
        )

    return _embeddings


# ============================================================
# FILE ENCODING
# ============================================================

def detect_file_encoding(file_path: str) -> str:
    """
    Detect text file encoding.
    """

    with open(file_path, "rb") as f:
        raw_data = f.read(50000)

    result = chardet.detect(raw_data)

    encoding = result.get("encoding")

    if not encoding:
        return "utf-8"

    return encoding


# ============================================================
# PDF
# ============================================================

def load_pdf(file_path: str):
    """
    Extract text from PDF.
    """

    reader = PdfReader(file_path)

    if reader.is_encrypted:
        raise ValueError("This PDF is encrypted and cannot be processed.")

    documents = []

    for page_number, page in enumerate(reader.pages, start=1):

        text = page.extract_text() or ""

        if text.strip():

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": os.path.basename(file_path),
                        "file_type": ".pdf",
                        "page": page_number,
                    },
                )
            )

    return documents


# ============================================================
# DOCX
# ============================================================

def load_docx(file_path: str):
    """
    Extract text from Microsoft Word DOCX.
    """

    text = docx_process(file_path)

    if not text.strip():
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source": os.path.basename(file_path),
                "file_type": ".docx",
            },
        )
    ]


# ============================================================
# TXT / MARKDOWN
# ============================================================

def load_text(file_path: str):
    """
    Load TXT or Markdown files.
    """

    encoding = detect_file_encoding(file_path)

    try:
        with open(file_path, "r", encoding=encoding, errors="replace") as f:
            text = f.read()

    except Exception:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()

    if not text.strip():
        return []

    extension = os.path.splitext(file_path)[1].lower()

    return [
        Document(
            page_content=text,
            metadata={
                "source": os.path.basename(file_path),
                "file_type": extension,
            },
        )
    ]


# ============================================================
# CSV
# ============================================================

# ============================================================
# CSV
# ============================================================

def load_csv(file_path: str):
    """
    Load CSV and group multiple rows into each Document.

    Grouping rows reduces the number of embeddings required
    for large CSV files and makes document processing faster.
    """

    encoding = detect_file_encoding(file_path)

    try:
        dataframe = pd.read_csv(
            file_path,
            encoding=encoding
        )

    except Exception:
        dataframe = pd.read_csv(
            file_path,
            encoding="utf-8",
            encoding_errors="replace"
        )

    if dataframe.empty:
        return []

    documents = []

    # Number of CSV rows stored in one Document
    ROWS_PER_DOCUMENT = 50

    for start in range(
        0,
        len(dataframe),
        ROWS_PER_DOCUMENT
    ):

        batch = dataframe.iloc[
            start:start + ROWS_PER_DOCUMENT
        ]

        rows_text = []

        for index, row in batch.iterrows():

            row_data = []

            for column in dataframe.columns:

                value = row[column]

                if pd.isna(value):
                    value = ""

                row_data.append(
                    f"{column}: {value}"
                )

            row_text = (
                f"Row {int(index) + 1}: "
                + " | ".join(row_data)
            )

            rows_text.append(row_text)

        # Include column names so the embedding has context
        columns_text = " | ".join(
            str(column)
            for column in dataframe.columns
        )

        document_text = (
            f"CSV Columns: {columns_text}\n\n"
            + "\n".join(rows_text)
        )

        documents.append(
            Document(
                page_content=document_text,
                metadata={
                    "source": os.path.basename(file_path),
                    "file_type": ".csv",
                    "row_start": int(start) + 1,
                    "row_end": int(start + len(batch)),
                    "rows_in_document": int(len(batch)),
                },
            )
        )

    return documents
# ============================================================
# EXCEL
# ============================================================

def load_excel(file_path: str):
    """
    Load XLSX or XLS files.

    Each spreadsheet row becomes a Document.
    Sheet name is preserved in metadata.
    """

    extension = os.path.splitext(file_path)[1].lower()

    engine = None

    if extension == ".xlsx":
        engine = "openpyxl"

    elif extension == ".xls":
        engine = "xlrd"

    excel_file = pd.ExcelFile(
        file_path,
        engine=engine
    )

    documents = []

    for sheet_name in excel_file.sheet_names:

        dataframe = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
            engine=engine
        )

        if dataframe.empty:
            continue

        for index, row in dataframe.iterrows():

            row_data = []

            for column in dataframe.columns:

                value = row[column]

                if pd.isna(value):
                    value = ""

                row_data.append(
                    f"{column}: {value}"
                )

            text = "\n".join(row_data)

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": os.path.basename(file_path),
                        "file_type": extension,
                        "sheet": str(sheet_name),
                        "row": int(index) + 1,
                    },
                )
            )

    return documents


# ============================================================
# JSON
# ============================================================

def load_json(file_path: str):
    """
    Load JSON and convert it into readable text.
    """

    encoding = detect_file_encoding(file_path)

    try:
        with open(
            file_path,
            "r",
            encoding=encoding,
            errors="replace"
        ) as f:
            data = json.load(f)

    except Exception as e:
        raise ValueError(f"Invalid JSON file: {e}")

    text = json.dumps(
        data,
        indent=2,
        ensure_ascii=False
    )

    if not text.strip():
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source": os.path.basename(file_path),
                "file_type": ".json",
            },
        )
    ]


# ============================================================
# POWERPOINT
# ============================================================

def load_pptx(file_path: str):
    """
    Extract text from PowerPoint slides.
    """

    presentation = Presentation(file_path)

    documents = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        combined_text = "\n".join(slide_text)

        if combined_text.strip():

            documents.append(
                Document(
                    page_content=combined_text,
                    metadata={
                        "source": os.path.basename(file_path),
                        "file_type": ".pptx",
                        "slide": slide_number,
                    },
                )
            )

    return documents


# ============================================================
# UNIVERSAL DOCUMENT LOADER
# ============================================================

def load_document(file_path: str):
    
    """
    Automatically detect the extension and load the document.
    """

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File does not exist: {file_path}"
        )

    extension = os.path.splitext(file_path)[1].lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}. "
            f"Supported types: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    if extension == ".pdf":
        return load_pdf(file_path)

    if extension == ".docx":
        return load_docx(file_path)

    if extension in [".txt", ".md"]:
        return load_text(file_path)

    if extension == ".csv":
        return load_csv(file_path)

    if extension in [".xlsx", ".xls"]:
        return load_excel(file_path)

    if extension == ".json":
        return load_json(file_path)

    if extension == ".pptx":
        return load_pptx(file_path)

    raise ValueError(
        f"No loader implemented for {extension}"
    )


# ============================================================
# CHUNKING
# ============================================================

def split_documents(documents):
    """
    Split documents into chunks.

    CSV documents are already grouped into batches,
    so use a larger chunk size for CSV files.
    """

    normal_splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=100,
        separators=[
            "\n\n",
            "\n",
            ". ",
            " ",
            ""
        ]
    )

    csv_splitter = RecursiveCharacterTextSplitter(
        chunk_size=1500,
        chunk_overlap=100,
        separators=[
            "\n",
            " | ",
            " ",
            ""
        ]
    )

    normal_documents = []
    csv_documents = []

    for document in documents:

        if document.metadata.get("file_type") == ".csv":
            csv_documents.append(document)
        else:
            normal_documents.append(document)

    chunks = []

    if normal_documents:
        chunks.extend(
            normal_splitter.split_documents(
                normal_documents
            )
        )

    if csv_documents:
        chunks.extend(
            csv_splitter.split_documents(
                csv_documents
            )
        )

    return chunks


# ============================================================
# USER-SPECIFIC CHROMA COLLECTION
# ============================================================

def get_collection_name(session_id: str) -> str:
    """
    Generate a safe Chroma collection name for each user/session.
    """

    clean_id = "".join(
        character
        for character in session_id
        if character.isalnum()
    )

    if not clean_id:
        clean_id = "default"

    collection_name = f"rag_{clean_id}"

    return collection_name[:63]


@lru_cache(maxsize=100)
def get_vector_store(session_id: str):
    """
    Get and cache the Chroma vector store for a session.
    """

    os.makedirs(
        CHROMA_PATH,
        exist_ok=True
    )

    collection_name = get_collection_name(session_id)

    return Chroma(
        collection_name=collection_name,
        persist_directory=CHROMA_PATH,
        embedding_function=get_embeddings(),
    )


# ============================================================
# INDEX DOCUMENT
# ============================================================

def index_document(file_path: str, session_id: str):
    """
    Complete ingestion pipeline with timing information.
    """

    total_start = time.perf_counter()

    # --------------------------------------------------------
    # 1. LOAD
    # --------------------------------------------------------

    start = time.perf_counter()

    documents = load_document(file_path)

    load_time = time.perf_counter() - start

    if not documents:
        raise ValueError(
            "No readable text was found in the document."
        )

    # --------------------------------------------------------
    # 2. CHUNK
    # --------------------------------------------------------

    start = time.perf_counter()

    chunks = split_documents(documents)

    chunk_time = time.perf_counter() - start

    if not chunks:
        raise ValueError(
            "The document could not be divided into chunks."
        )

    # --------------------------------------------------------
    # 3. VECTOR STORE / EMBEDDINGS
    # --------------------------------------------------------

    start = time.perf_counter()

    vector_store = get_vector_store(session_id)

    vector_store.add_documents(chunks)

    embedding_time = time.perf_counter() - start

    # --------------------------------------------------------
    # TOTAL
    # --------------------------------------------------------

    total_time = time.perf_counter() - total_start

    print("\n" + "=" * 60)
    print("DOCUMENT PROCESSING COMPLETE")
    print("=" * 60)

    print(f"Documents loaded : {len(documents)}")
    print(f"Chunks created   : {len(chunks)}")
    print(f"Loading time     : {load_time:.2f}s")
    print(f"Chunking time    : {chunk_time:.2f}s")
    print(f"Embedding/index  : {embedding_time:.2f}s")
    print(f"TOTAL TIME       : {total_time:.2f}s")

    print("=" * 60 + "\n")

    return {
        "documents": len(documents),
        "chunks": len(chunks),
        "collection": get_collection_name(session_id),
        "load_time": round(load_time, 2),
        "chunk_time": round(chunk_time, 2),
        "embedding_time": round(embedding_time, 2),
        "total_time": round(total_time, 2),
    }
# ============================================================
# RETRIEVE
# ============================================================

def retrieve_documents(
    question: str,
    session_id: str,
    k: int = 5
):
    """
    Retrieve the most relevant chunks for a question.
    """

    vector_store = get_vector_store(session_id)

    try:
        documents = vector_store.similarity_search(
            question,
            k=k
        )
    except Exception:
        documents = []

    return documents


# ============================================================
# FORMAT RETRIEVED CONTEXT
# ============================================================

def format_context(documents):
    """
    Convert retrieved documents into context for the LLM.
    """

    if not documents:
        return "No relevant information was found."

    formatted = []

    for index, document in enumerate(
        documents,
        start=1
    ):

        metadata = document.metadata

        source = metadata.get(
            "source",
            "Unknown"
        )

        location_parts = []

        if "page" in metadata:
            location_parts.append(
                f"page {metadata['page']}"
            )

        if "row" in metadata:
            location_parts.append(
                f"row {metadata['row']}"
            )

        if "sheet" in metadata:
            location_parts.append(
                f"sheet {metadata['sheet']}"
            )

        if "slide" in metadata:
            location_parts.append(
                f"slide {metadata['slide']}"
            )

        location = ""

        if location_parts:
            location = (
                " (" +
                ", ".join(location_parts) +
                ")"
            )

        formatted.append(
            f"--- Source {index}: "
            f"{source}{location} ---\n"
            f"{document.page_content}"
        )

    return "\n\n".join(formatted)
